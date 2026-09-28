#!/usr/bin/env python3
"""Build a Chinese Word guide for the Nature-style candidate figure suite."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


FIGURES = [
    ("图 1", "分析流程与已完成工作状态", "figure_1_workflow_and_status.png", "流程图将正式 HMM scan、strict tier1 核心家族并集、PHA 证据分层、古菌 PhaZh1-like audit、domain annotation 和跨门人工复核串成一条可审计链路。右侧输出规模以 log10 记录数表示，便于同时展示从千级任务到万级候选的数量级变化。", "核心结论：项目的主要计算阶段已经形成连续、可追溯的候选分析流程。图中 status 仅表示计算阶段完成，不表示实验验证完成。"),
    ("图 2", "PHA 候选蛋白的亚型与证据层分布", "figure_2_pha_subtype_distribution.png", "该图统计 109,087 条去重候选蛋白，并按 ePhaZ/iPhaZ 及其 secreted、core、tier2 review 等证据层分组。条形长度是候选蛋白数，分层名称保留了项目中的机器可审计分类。", "核心结论：候选集合主要由 iPhaZ intracellular、ePhaZ tier2 review 和 ePhaZ SCL secreted 等层构成。这里的 subtype 是证据分层，不是实验证实的酶活类型。"),
    ("图 3", "Strict tier1 核心家族的序列与基因组规模", "figure_3_core_family_scale.png", "左、中央面板分别比较四个核心家族的 tier1 序列数和覆盖基因组数；右侧面板汇总四家族 strict union（38,741 个基因组），并标出至少包含两个家族的 6,578 个基因组。大柱标签采用柱内白字以避免与标题冲突。", "核心结论：iPhaZ 和 ArchPhaZ hydrolase 在候选规模上占主导，四家族交集提供了更严格的候选基因组层。该层仍是高置信计算候选，不等于 PHB/PHA 降解表型。"),
    ("图 4", "古菌候选的门水平分布", "figure_4_archaeal_phylum_distribution.png", "左面板显示 2,307 条古菌 audit 记录在 Halobacteriota、Thermoproteota 和 Thermoplasmatota 中的分层组成；右面板单独放大 PhaZh1-like high/review 优先层，避免 exploratory 记录淹没重点。", "核心结论：high/review 候选主要落在 Halobacteriota，少量分布于 Thermoproteota。该图只代表拥有实际 taxonomy 的古菌 audit，不能外推到全部 PHA 候选蛋白。"),
    ("图 5", "古菌 audit 分层与 N 端 GxSxG motif", "figure_5_archaeal_audit_and_motif.png", "左面板给出全量古菌 audit 的四层记录数：PhaZh1-like high 197、review 7、patatin exploratory 2,027、non-PhaZh1 patatin 76。右面板统计 domain-annotated high/review 候选中 N 端 GxSxG motif 的有无。", "核心结论：197 条 high 和 7 条 review 记录构成优先复核层，且大多数 domain-annotated 记录带有 GxSxG motif。motif 是序列特征证据，不能独立证明催化活性。"),
    ("图 6", "物种代表性与跨门人工复核", "figure_6_species_and_manual_review.png", "左面板列出 high/review 记录中出现频率最高的物种标签；右面板展示 9 条人工跨门复核记录，以 candidate-HMM coverage 和 -log10(E-value) 表示相对支持强度，并用颜色/形状区分 high 与 review。", "核心结论：优先候选在物种层面呈现明显的盐生古菌聚集，同时保留了跨门 review 记录以控制分类边界风险。物种标签来自 audit taxonomy 字符串，人工复核仍是 candidate-only。"),
    ("图 7", "ArchPhaZ_patatin 阳性/阴性对照校准", "supplementary_figure_1_hmm_calibration.png", "左面板展示阳性和阴性对照的 domain bitscore 分布，右面板汇总命中数。阳性对照命中 8 条，阴性对照命中 5 条。", "核心结论：注册的 ArchPhaZ_patatin 模型能够命中阳性对照，但阴性 panel 也有命中，因此不能把该模型解释为零交叉反应的特异性判别器。"),
    ("图 8", "古菌候选邻域 marker 共现", "supplementary_figure_2_neighborhood_markers.png", "左面板按 audit layer 展示 BdhA、PhaJ、PhaC、PhaE 的邻域共现比例；右面板给出所有 audit 记录合并后的总体比例。数值是包含该 marker 的记录百分比。", "核心结论：邻域 marker 可作为局部基因组上下文的辅助证据，帮助安排人工复核优先级，但不能据此证明通路归属、共转录或 PHB 降解。"),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, name="Calibri", size=11, color="000000", bold=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def add_paragraph(doc, text="", style=None, before=0, after=6, line=1.25, align=None):
    p = doc.add_paragraph(style=style)
    if text:
        run = p.add_run(text)
        set_run_font(run)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if align is not None:
        p.alignment = align
    return p


def build(run_dir: Path, output: Path) -> None:
    fig_dir = run_dir / "results" / "figures"
    if not fig_dir.is_dir():
        raise FileNotFoundError(fig_dir)
    missing = [name for _, _, name, _, _ in FIGURES if not (fig_dir / name).is_file()]
    if missing:
        raise FileNotFoundError("missing figure previews: " + ", ".join(missing))

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for style_name, size, color, before, after in (("Heading 1", 16, "2E74B5", 18, 10), ("Heading 2", 13, "2E74B5", 14, 7), ("Heading 3", 12, "1F4D78", 10, 5)):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run("PHB/PHA 候选分析 Nature 风格图组说明")
    set_run_font(r, size=22, color="1F4D78", bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("6 张主图 + 2 张补充图 | dated run: 20260906_nature_figure_suite_03")
    set_run_font(sr, size=10, color="767676")
    add_paragraph(doc, "本文件依次说明图组所展示的分析流程、候选规模、古菌 taxonomy 分布、domain/motif/邻域证据和人工复核结果。所有图均来自同一 dated run 的输入快照与 source-data 表。", after=8)

    callout = doc.add_table(rows=1, cols=1)
    callout.alignment = WD_TABLE_ALIGNMENT.LEFT
    callout.autofit = False
    callout.columns[0].width = Inches(6.5)
    cell = callout.cell(0, 0)
    cell.width = Inches(6.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, "F4F6F9")
    set_cell_margins(cell, top=120, bottom=120)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    rr = p.add_run("统一证据边界：")
    set_run_font(rr, size=10, color="7A5A00", bold=True)
    rr = p.add_run("HMM、domain、motif、邻域、SignalP 和 taxonomy 结果只能支持候选同源或功能潜力，不等同于已验证的 PHB/PHA 降解表型。")
    set_run_font(rr, size=10, color="000000")

    doc.add_heading("图组总览", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    widths = [Inches(0.65), Inches(2.6), Inches(3.25)]
    headers = ["编号", "主题", "主要数据口径"]
    for i, (cell, text, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
        cell.width = width
        set_cell_margins(cell)
        set_cell_shading(cell, "E8EEF5")
        run = cell.paragraphs[0].add_run(text)
        set_run_font(run, size=9, color="1F4D78", bold=True)
    for label, title_text, _, _, _ in FIGURES:
        cells = table.add_row().cells
        for cell, width in zip(cells, widths):
            cell.width = width
            set_cell_margins(cell)
        for cell, text in zip(cells, [label, title_text, "candidate-only; 见对应图注"]):
            run = cell.paragraphs[0].add_run(text)
            set_run_font(run, size=9)

    for idx, (label, title_text, filename, explanation, conclusion) in enumerate(FIGURES):
        doc.add_page_break()
        doc.add_heading(f"{label}｜{title_text}", level=1)
        pic = doc.add_picture(str(fig_dir / filename), width=Inches(6.45))
        last = doc.paragraphs[-1]
        last.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(8)
        cr = cap.add_run(f"{label}  {title_text}")
        set_run_font(cr, size=9, color="767676", bold=True)
        h = doc.add_paragraph()
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run("图中展示：")
        set_run_font(r, size=11, color="1F4D78", bold=True)
        r = h.add_run(explanation)
        set_run_font(r)
        c = doc.add_paragraph()
        c.paragraph_format.space_after = Pt(6)
        r = c.add_run("解读：")
        set_run_font(r, size=11, color="1F4D78", bold=True)
        r = c.add_run(conclusion)
        set_run_font(r)

    doc.add_page_break()
    doc.add_heading("数据与复现信息", level=1)
    add_paragraph(doc, "图像文件、source-data TSV、输入契约和导出哈希均保存在同一 dated run 下。图组使用 Python/Matplotlib 生成；SVG 保留可编辑文字，PDF 使用 TrueType 字体，TIFF 为 600 dpi。", after=6)
    add_paragraph(doc, "主要规模核对：PHA 亚型候选蛋白 109,087 条；古菌 audit 记录 2,307 条，其中 PhaZh1-like high 197 条、review 7 条；跨门人工复核 9 条；阳性/阴性校准命中分别为 8/5。", after=6)
    add_paragraph(doc, "建议引用：将本说明文档与 figure_manifest.json、input_contract.json 及 11 份 source-data TSV 一并归档。正式论文中应在图注中明确 n 的定义、数据来源和 candidate-only 证据边界。", after=6)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("PHB/PHA candidate-analysis figure suite | candidate-only")
    set_run_font(fr, size=8, color="767676")
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.run_dir, args.output)
    print(f"Wrote {args.output}")


if __name__ == "__main__":
    main()
